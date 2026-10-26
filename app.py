import streamlit as st
import time
from src.agents.agents import build_search_agent, build_reader_agent, writer_chain, critic_chain

# ── Page config ──────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Multi-Agent Research Assistant",
    page_icon="🔬",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ── Custom CSS ────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Manrope:wght@400;500;600;700;800&display=swap');
:root { --ink:#edf3f8; --muted:#a7b6c7; --line:#2c3b4b; --accent:#83dbc7; }
html,body,[data-testid="stApp"] {font-family:'Manrope',sans-serif;color:var(--ink);}
.stApp {background:#0d151f;}
[data-testid="stHeader"] {background:transparent;}
#MainMenu,footer {visibility:hidden;}
.block-container {max-width:1280px;padding:2rem 3rem 4rem;}
h1,h2,h3 {color:var(--ink);font-family:'Manrope',sans-serif;letter-spacing:-.035em;}
.hero {text-align:left;padding:3rem 0 1rem;max-width:800px;}
.hero-eyebrow {font-size:.7rem;text-transform:uppercase;letter-spacing:.18em;color:var(--accent);font-weight:700;margin-bottom:1.2rem;}
.hero h1 {font-family:'Manrope',sans-serif;font-size:clamp(2.4rem,5vw,4rem);line-height:1.12;letter-spacing:-.055em;font-weight:800;margin:0 0 1.2rem;}
.hero h1 span {color:var(--accent);}
.hero-sub {font-size:1rem;line-height:1.8;color:var(--muted);max-width:620px;}
.divider {height:1px;background:var(--line);margin:1.8rem 0 2rem;}
[data-testid="stVerticalBlockBorderWrapper"] > div {background:#141f2b;border-color:var(--line)!important;border-radius:16px!important;}
[data-testid="stVerticalBlockBorderWrapper"] > div > [data-testid="stVerticalBlock"] {padding:.6rem;}
.section-heading {font-size:1.25rem;letter-spacing:-.03em;font-weight:700;margin:0 0 1rem;color:var(--ink);}
.card-copy {font-size:.87rem;color:var(--muted);line-height:1.7;margin-bottom:1.5rem;}
[data-testid="stWidgetLabel"] p {color:var(--ink);font-size:.82rem;font-weight:600;}
[data-testid="stTextInput"] [data-baseweb="input"] {background:#0c1722;border:1px solid #405268;border-radius:10px;}
[data-testid="stTextInput"] input {background:#0c1722!important;color:var(--ink)!important;min-height:50px;}
[data-testid="stTextInput"] input::placeholder {color:#9aaabd;opacity:1;}
[data-testid="stTextInput"] [data-baseweb="input"]:focus-within {border-color:var(--accent);box-shadow:0 0 0 3px #83dbc720;}
.stButton > button {background:var(--accent)!important;border:1px solid var(--accent)!important;color:#102a25!important;border-radius:10px;font-weight:700;min-height:49px;width:100%;}
.stButton > button:hover {background:#a0e7d7!important;}
button:focus-visible {outline:3px solid #a0e7d7!important;outline-offset:3px;}
.step-card {background:#111c28;border:1px solid var(--line);border-radius:12px;padding:1.1rem;margin-bottom:.75rem;}
.step-card.active {border-color:var(--accent);background:#19332f;}
.step-card.done {border-color:#365d56;background:#172b29;}
.step-header {display:flex;align-items:center;gap:.65rem;flex-wrap:wrap;}
.step-num {font-size:.7rem;color:var(--muted);}
.step-title {font-size:.88rem;font-weight:700;color:var(--ink);}
.step-status {margin-left:auto;font-size:.6rem;letter-spacing:.06em;font-weight:700;}
.status-waiting {color:var(--muted);}
.status-running,.status-done {color:var(--accent);}
.panel-label,.result-panel-title {font-size:.7rem;letter-spacing:.12em;text-transform:uppercase;font-weight:700;color:var(--accent);padding-bottom:1rem;border-bottom:1px solid var(--line);margin-bottom:1rem;}
.result-content {font-size:.9rem;line-height:1.8;color:#dae4ee;white-space:pre-wrap;overflow-wrap:anywhere;}
[data-testid="stMarkdownContainer"] {color:var(--ink);}
[data-testid="stMarkdownContainer"] p,[data-testid="stMarkdownContainer"] li {line-height:1.8;}
[data-testid="stMarkdownContainer"] a {color:#a0e7d7;}
[data-testid="stExpander"] {background:#111c28;border-color:var(--line);border-radius:10px;}
[data-testid="stExpander"] summary {color:var(--muted);}
.stDownloadButton > button {background:#1c3531;color:#a0e7d7;border:1px solid #456c63;border-radius:9px;min-height:43px;}
.notice {font-size:.72rem;color:var(--muted);border-top:1px solid var(--line);padding-top:1.4rem;margin-top:3rem;}
@media(max-width:760px) {.block-container {padding:1rem 1rem 2rem;}.hero {padding:2rem 0 1rem;}.step-card {padding:.9rem;}}
</style>
""", unsafe_allow_html=True)

# ── Document exports ─────────────────────────────────────────────────────────
@st.cache_data(show_spinner=False)
def build_document_downloads(report: str):
    """Convert the generated Markdown report to in-memory PDF and DOCX files."""
    from io import BytesIO
    from pathlib import Path
    from xml.sax.saxutils import escape

    import markdown
    import reportlab
    from bs4 import BeautifulSoup
    from docx import Document
    from docx.shared import Inches, Pt
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    )

    soup = BeautifulSoup(
        markdown.markdown(report, extensions=["tables", "fenced_code", "sane_lists"]),
        "html.parser",
    )
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    # Prefer a Unicode font, with a portable fallback bundled with ReportLab.
    font_candidates = [
        (Path("/usr/share/fonts/truetype/dejavu"), "DejaVuSans.ttf", "DejaVuSans-Bold.ttf"),
        (Path("C:/Windows/Fonts"), "arial.ttf", "arialbd.ttf"),
        (Path(reportlab.__file__).parent / "fonts", "Vera.ttf", "VeraBd.ttf"),
    ]
    for directory, regular, bold in font_candidates:
        if (directory / regular).exists() and (directory / bold).exists():
            pdfmetrics.registerFont(TTFont("ResearchBody", str(directory / regular)))
            pdfmetrics.registerFont(TTFont("ResearchBold", str(directory / bold)))
            break
    else:
        raise RuntimeError("No PDF font available. Install DejaVu Sans to enable PDF exports.")
    pdfmetrics.registerFontFamily(
        "ResearchBody", normal="ResearchBody", bold="ResearchBold",
        italic="ResearchBody", boldItalic="ResearchBold",
    )
    body_style = ParagraphStyle(
        "ResearchBodyStyle", fontName="ResearchBody", fontSize=10,
        leading=16, spaceAfter=9, textColor=colors.HexColor("#243344"),
        alignment=TA_LEFT, splitLongWords=True,
    )
    heading_styles = {
        level: ParagraphStyle(
            f"ResearchHeading{level}", parent=body_style,
            fontName="ResearchBold", fontSize={1: 21, 2: 16, 3: 13}.get(level, 11),
            leading={1: 27, 2: 22, 3: 19}.get(level, 17),
            spaceBefore=14, spaceAfter=8, keepWithNext=True,
        )
        for level in range(1, 7)
    }
    cell_style = ParagraphStyle(
        "ResearchCell", parent=body_style, fontSize=8, leading=11, spaceAfter=0,
    )
    story = []

    def inline_pdf(node):
        if isinstance(node, str):
            return escape(str(node))
        content = "".join(inline_pdf(child) for child in node.children)
        if node.name in ("strong", "b"):
            return f"<b>{content}</b>"
        if node.name in ("em", "i"):
            return f"<i>{content}</i>"
        if node.name == "br":
            return "<br/>"
        if node.name == "a":
            href = node.get("href", "")
            if href.startswith(("https://", "http://", "mailto:")):
                return f'<link href="{escape(href, {chr(34): "&quot;"})}" color="#24796c">{content}</link>'
        if node.name == "img":
            return escape(node.get("alt", ""))
        return content

    def inline_word(paragraph, node, bold=False, italic=False):
        if isinstance(node, str):
            run = paragraph.add_run(str(node))
            run.bold, run.italic = bold, italic
            return
        if node.name == "br":
            paragraph.add_run().add_break()
            return
        if node.name == "img":
            paragraph.add_run(node.get("alt", ""))
            return
        for child in node.children:
            inline_word(paragraph, child, bold or node.name in ("b", "strong"),
                        italic or node.name in ("i", "em"))
        if node.name == "a" and node.get("href", "").startswith(("http://", "https://")):
            paragraph.add_run(f" ({node['href']})")

    def render_block(node, prefix="", list_style=None):
        if isinstance(node, str):
            if node.strip():
                document.add_paragraph(str(node))
                story.append(Paragraph(escape(str(node)), body_style))
            return
        name = node.name
        if name in ("ul", "ol"):
            for index, item in enumerate(node.find_all("li", recursive=False), 1):
                render_block(item, f"{index}. " if name == "ol" else "• ",
                             "List Number" if name == "ol" else "List Bullet")
        elif name == "table":
            rows = [row.find_all(["th", "td"], recursive=False) for row in node.find_all("tr")]
            rows = [row for row in rows if row]
            if not rows:
                return
            count = max(map(len, rows))
            table = document.add_table(rows=0, cols=count)
            table.style = "Table Grid"
            pdf_rows = []
            for row in rows:
                cells = table.add_row().cells
                pdf_row = []
                for index in range(count):
                    cell = row[index] if index < len(row) else None
                    if cell is not None:
                        inline_word(cells[index].paragraphs[0], cell, bold=cell.name == "th")
                        pdf_row.append(Paragraph(inline_pdf(cell) or " ", cell_style))
                    else:
                        pdf_row.append(Paragraph(" ", cell_style))
                pdf_rows.append(pdf_row)
            pdf_table = Table(pdf_rows, colWidths=[468 / count] * count,
                              repeatRows=1 if rows[0][0].name == "th" else 0,
                              splitInRow=1, hAlign="LEFT")
            pdf_table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5df")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#edf5f3")),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.extend([pdf_table, Spacer(1, 12)])
        elif name == "pre":
            text = node.get_text()
            paragraph = document.add_paragraph()
            run = paragraph.add_run(text)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            # Separate wrapped lines can flow across PDF pages, even for long code blocks.
            for line in text.splitlines():
                for wrapped in simpleSplit(line or " ", "ResearchBody", 9, 460):
                    story.append(Paragraph(escape(wrapped) or " ", body_style))
        elif name in ("div", "blockquote"):
            for child in node.children:
                render_block(child)
        elif name == "hr":
            story.append(Spacer(1, 12))
        else:
            heading = int(name[1]) if name in ("h1", "h2", "h3", "h4", "h5", "h6") else None
            paragraph = (document.add_heading(level=heading) if heading
                         else document.add_paragraph(style=list_style))
            # Nested lists are handled separately to preserve their structure.
            content = []
            for child in node.children:
                if getattr(child, "name", None) not in ("ul", "ol"):
                    inline_word(paragraph, child)
                    content.append(inline_pdf(child))
            story.append(Paragraph(escape(prefix) + "".join(content) or " ",
                                   heading_styles[heading] if heading else body_style))
            for child in node.find_all(["ul", "ol"], recursive=False):
                render_block(child)

    for block in soup.children:
        render_block(block)

    word_buffer = BytesIO()
    document.save(word_buffer)
    pdf_buffer = BytesIO()
    pdf_document = SimpleDocTemplate(
        pdf_buffer, pagesize=(612, 792), rightMargin=72, leftMargin=72,
        topMargin=58, bottomMargin=58, title="Research Report",
    )
    pdf_document.build(story or [Paragraph("Research Report", body_style)])
    return pdf_buffer.getvalue(), word_buffer.getvalue()

# ── Helper: render a step card ────────────────────────────────────────────────
def step_card(num: str, title: str, state: str, desc: str = ""):
    status_map = {
        "waiting": ("WAITING", "status-waiting"),
        "running": ("● RUNNING", "status-running"),
        "done":    ("✓ DONE",   "status-done"),
    }

    label, cls = status_map.get(state, ("", ""))
    card_cls = {"running": "active", "done": "done"}.get(state, "")

    st.markdown(f"""
    <div class="step-card {card_cls}">
        <div class="step-header">
            <span class="step-num">{num}</span>
            <span class="step-title">{title}</span>
            <span class="step-status {cls}">{label}</span>
        </div>
        {"<div style='font-size:0.82rem;color:#94a3b8;margin-top:0.3rem;'>"+desc+"</div>" if desc else ""}
    </div>
    """, unsafe_allow_html=True)


# ── Session state init ────────────────────────────────────────────────────────
for key in ("results", "running", "done"):
    if key not in st.session_state:
        st.session_state[key] = {} if key == "results" else False


# ── Hero ──────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="hero">
    <div class="hero-eyebrow">ResearchAgent / Your research workspace</div>
    <h1>Deep research.<br><span>Thoughtfully composed.</span></h1>
    <p class="hero-sub">
        Four specialized AI agents collaborate — searching, scraping, writing,
        and critiquing — to deliver a polished research report on any topic.
    </p>
</div>

<div class="divider"></div>
""", unsafe_allow_html=True)


# ── Layout: input left, pipeline right ───────────────────────────────────────
col_input, col_pipeline = st.columns([1.3, 1], gap="large")

with col_input, st.container(border=True):

    st.markdown('<div class="section-heading">Define your research</div><div class="card-copy">Start with a topic or a specific question. A clear brief gives your research direction.</div>', unsafe_allow_html=True)

    topic = st.text_input(
        "Research Topic",
        placeholder="e.g. Roadmap for AGI development in next 5 years",
        key="topic_input",
        label_visibility="visible",
    )

    run_btn = st.button(
        "Start research →",
        use_container_width=True
    )



    st.markdown("""
    <div style="margin-top:1.5rem;color:#a7b6c7;font-size:.7rem;letter-spacing:.12em;text-transform:uppercase;">Ideas to explore</div>
    <div style="display:flex;gap:.5rem;flex-wrap:wrap;margin-top:.7rem;">
        <span style="background:#111c28;border:1px solid #2c3b4b;border-radius:7px;padding:.5rem .7rem;font-size:.75rem;color:#c8d5e2;">The future of LLMs in tech</span>
        <span style="background:#111c28;border:1px solid #2c3b4b;border-radius:7px;padding:.5rem .7rem;font-size:.75rem;color:#c8d5e2;">AI agents in 2026</span>
        <span style="background:#111c28;border:1px solid #2c3b4b;border-radius:7px;padding:.5rem .7rem;font-size:.75rem;color:#c8d5e2;">The roadmap to AGI</span>
    </div>
    """, unsafe_allow_html=True)
with col_pipeline, st.container(border=True):

    st.markdown(
        '<div class="section-heading">Agent workflow</div>',
        unsafe_allow_html=True
    )

    # Each card has a persistent slot that can be replaced during this same run.
    workflow_steps = [
        ("search", "01", "Search Agent", "Gathers recent web information"),
        ("reader", "02", "Reader Agent", "Scrapes & extracts deep content"),
        ("writer", "03", "Writer Chain", "Drafts the full research report"),
        ("critic", "04", "Critic Chain", "Reviews & scores the report"),
    ]
    workflow_slots = {key: st.empty() for key, _, _, _ in workflow_steps}

    def s(step):
        # Read the current session state, rather than an earlier results snapshot.
        current_results = st.session_state.results
        if step in current_results:
            return "done"
        if st.session_state.running:
            for key, _, _, _ in workflow_steps:
                if key not in current_results:
                    return "running" if key == step else "waiting"
        return "waiting"

    def render_workflow():
        for key, number, title, description in workflow_steps:
            with workflow_slots[key]:
                step_card(number, title, s(key), description)

    render_workflow()

# ── Run pipeline ──────────────────────────────────────────────────────────────
if run_btn:

    if not topic.strip():
        st.warning("Please enter a research topic first.")

    else:
        st.session_state.results = {}
        st.session_state.running = True
        st.session_state.done = False
        st.rerun()


if st.session_state.running and not st.session_state.done:

    results = {}
    topic_val = st.session_state.topic_input

    # ── Step 1: Search ──
    with st.spinner("🔍 Search Agent is working…"):

        search_agent = build_search_agent()

        sr = search_agent.invoke({
            "messages": [
                ("user",
                 f"Find recent, reliable and detailed information about: {topic_val}")
            ]
        })

        results["search"] = sr["messages"][-1].content
        st.session_state.results = dict(results)
        render_workflow()

    # ── Step 2: Reader ──
    with st.spinner("📄 Reader Agent is scraping top resources…"):

        reader_agent = build_reader_agent()

        rr = reader_agent.invoke({
            "messages": [(
                "user",
                f"Based on the following search results about '{topic_val}', "
                f"pick the most relevant URL and scrape it for deeper content.\n\n"
                f"Search Results:\n{results['search'][:800]}"
            )]
        })

        results["reader"] = rr["messages"][-1].content
        st.session_state.results = dict(results)
        render_workflow()

    # ── Step 3: Writer ──
    with st.spinner("✍️ Writer is drafting the report…"):

        research_combined = (
            f"SEARCH RESULTS:\n{results['search']}\n\n"
            f"DETAILED SCRAPED CONTENT:\n{results['reader']}"
        )

        results["writer"] = writer_chain.invoke({
            "topic": topic_val,
            "research": research_combined
        })

        st.session_state.results = dict(results)
        render_workflow()

    # ── Step 4: Critic ──
    with st.spinner("🧐 Critic is reviewing the report…"):

        results["critic"] = critic_chain.invoke({
            "report": results["writer"]
        })

        st.session_state.results = dict(results)
        render_workflow()

    st.session_state.running = False
    st.session_state.done = True

    st.rerun()


# ── Results display ───────────────────────────────────────────────────────────
r = st.session_state.results

if r:

    st.markdown('<div class="divider"></div>', unsafe_allow_html=True)

    st.markdown(
        '<div class="section-heading">Your research findings</div>',
        unsafe_allow_html=True
    )

    # Raw outputs
    if "search" in r:
        with st.expander("🔍 Search Results (raw)", expanded=False):

            st.markdown(
                f'''
                <div class="result-panel">
                    <div class="result-panel-title">
                        Search Agent Output
                    </div>

                    <div class="result-content">
                        {r["search"]}
                    </div>
                </div>
                ''',
                unsafe_allow_html=True
            )

    if "reader" in r:
        with st.expander("📄 Scraped Content (raw)", expanded=False):

            st.markdown(
                f'''
                <div class="result-panel">
                    <div class="result-panel-title">
                        Reader Agent Output
                    </div>

                    <div class="result-content">
                        {r["reader"]}
                    </div>
                </div>
                ''',
                unsafe_allow_html=True
            )

    # Final report
    if "writer" in r:

        with st.container(border=True):
            st.markdown('<div class="panel-label">Final research report</div>', unsafe_allow_html=True)
            st.markdown(r["writer"])

        # Keep Markdown available even when optional document packages are missing.
        filename = f"research_report_{int(time.time())}"
        md_col, pdf_col, word_col = st.columns(3)
        with md_col:
            st.download_button(
                label="Download Markdown", data=r["writer"],
                file_name=f"{filename}.md", mime="text/markdown",
                use_container_width=True,
            )
        try:
            pdf_data, word_data = build_document_downloads(str(r["writer"]))
        except ImportError:
            st.info("Enable PDF and Word downloads by installing: "
                    "pip install python-docx reportlab Markdown beautifulsoup4")
        except Exception:
            st.warning("PDF and Word exports could not be created for this report. "
                       "You can still download the Markdown version.")
        else:
            with pdf_col:
                st.download_button(
                    label="Download PDF", data=pdf_data,
                    file_name=f"{filename}.pdf", mime="application/pdf",
                    use_container_width=True,
                )
            with word_col:
                st.download_button(
                    label="Download Word", data=word_data,
                    file_name=f"{filename}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True,
                )

    # Critic feedback
    if "critic" in r:

        with st.container(border=True):
            st.markdown('<div class="panel-label">Review & feedback</div>', unsafe_allow_html=True)
            st.markdown(r["critic"])


# ── Footer ────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="notice">
    ResearchAgent · Powered by LangChain multi-agent pipeline · Built with Streamlit
</div>
""", unsafe_allow_html=True)